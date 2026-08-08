"""Excel Generator usando openpyxl"""

import logging
from pathlib import Path
from typing import Dict, Any
from datetime import datetime

logger = logging.getLogger(__name__)


class ExcelGenerator:
    """Genera documentos Excel profesionales"""

    def __init__(self, theme: Dict[str, Any], language: str = "es"):
        self.theme = theme
        self.language = language

    def generate(self, document_type: str, data: Dict[str, Any], output_path: Path) -> None:
        """Genera Excel"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            logger.error("openpyxl no instalado. Instala con: pip install openpyxl")
            return

        wb = Workbook()
        ws = wb.active
        ws.title = document_type

        # Estilos
        primary_color = self.theme.get("primary_color", "003366").lstrip("#")
        header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)

        # Agregar titulo
        ws["A1"] = f"REPORTE: {document_type.upper()}"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:D1")

        # Agregar metadata
        ws["A3"] = "Empresa:"
        ws["B3"] = data.get("empresa", "")
        ws["A4"] = "Generado:"
        ws["B4"] = datetime.now().strftime("%d/%m/%Y %H:%M")

        # Contenido según tipo
        if document_type == "invoice":
            self._generate_invoice_sheet(ws, data, header_fill, header_font)
        elif document_type == "report":
            self._generate_report_sheet(ws, data, header_fill, header_font)
        else:
            self._generate_generic_sheet(ws, data, header_fill, header_font)

        # Ajustar ancho de columnas
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

        wb.save(str(output_path))
        logger.info(f"Excel document saved: {output_path}")

    def _generate_invoice_sheet(self, ws, data, header_fill, header_font):
        """Genera hoja de factura"""
        ws["A6"] = "Descripción"
        ws["B6"] = "Precio Unitario"
        ws["C6"] = "Cantidad"
        ws["D6"] = "Total"

        for cell in ["A6", "B6", "C6", "D6"]:
            ws[cell].fill = header_fill
            ws[cell].font = header_font

        row = 7
        for item in data.get("items", []):
            ws[f"A{row}"] = item.get("descripcion", "")
            ws[f"B{row}"] = item.get("precio_unitario", 0)
            ws[f"C{row}"] = item.get("cantidad", 0)
            ws[f"D{row}"] = f"=B{row}*C{row}"
            row += 1

        # Totales
        row += 1
        ws[f"C{row}"] = "SUBTOTAL:"
        ws[f"D{row}"] = f"=SUM(D7:D{row-2})"
        row += 1
        ws[f"C{row}"] = "IMPUESTO (19%):"
        ws[f"D{row}"] = f"=D{row-1}*0.19"
        row += 1
        ws[f"C{row}"] = "TOTAL:"
        ws[f"D{row}"] = f"=D{row-2}+D{row-1}"
        ws[f"D{row}"].font = Font(bold=True, size=12)

    def _generate_report_sheet(self, ws, data, header_fill, header_font):
        """Genera hoja de reporte"""
        row = 6
        ws[f"A{row}"] = "Métrica"
        ws[f"B{row}"] = "Valor"

        for cell in [f"A{row}", f"B{row}"]:
            ws[cell].fill = header_fill
            ws[cell].font = header_font

        row += 1
        for key, value in data.get("metricas", {}).items():
            ws[f"A{row}"] = key
            ws[f"B{row}"] = value
            row += 1

    def _generate_generic_sheet(self, ws, data, header_fill, header_font):
        """Genera hoja genérica"""
        row = 6
        for key, value in data.items():
            if key not in ["empresa", "titulo"]:
                ws[f"A{row}"] = str(key).upper()
                ws[f"B{row}"] = str(value)
                row += 1
