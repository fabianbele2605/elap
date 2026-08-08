"""PowerPoint Generator usando python-pptx"""

import logging
from pathlib import Path
from typing import Dict, Any
from datetime import datetime

logger = logging.getLogger(__name__)


class PowerPointGenerator:
    """Genera presentaciones PowerPoint profesionales"""

    def __init__(self, theme: Dict[str, Any], language: str = "es"):
        self.theme = theme
        self.language = language

    def generate(self, document_type: str, data: Dict[str, Any], output_path: Path) -> None:
        """Genera PowerPoint"""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.enum.text import PP_ALIGN
            from pptx.dml.color import RGBColor
        except ImportError:
            logger.error("python-pptx no instalado. Instala con: pip install python-pptx")
            return

        # Crear presentación
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)

        # Colores del tema
        primary_color = self._hex_to_rgb(self.theme.get("primary_color", "#003366"))
        secondary_color = self._hex_to_rgb(self.theme.get("secondary_color", "#0099FF"))

        # Slide 1: Portada
        slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        background = slide1.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*primary_color)

        title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1.5))
        title_frame = title_box.text_frame
        title_frame.text = data.get("titulo", document_type.upper())
        title_frame.paragraphs[0].font.size = Pt(54)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

        subtitle_box = slide1.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(9), Inches(1))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = self.theme.get("name", "Empresa")
        subtitle_frame.paragraphs[0].font.size = Pt(28)
        subtitle_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

        # Slide 2: Contenido
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])
        title_box2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
        title_frame2 = title_box2.text_frame
        title_frame2.text = "Contenido"
        title_frame2.paragraphs[0].font.size = Pt(40)
        title_frame2.paragraphs[0].font.bold = True
        title_frame2.paragraphs[0].font.color.rgb = RGBColor(*primary_color)

        content_box = slide2.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
        content_frame = content_box.text_frame
        content_frame.word_wrap = True
        content_frame.text = data.get("contenido_ia", "Sin contenido")
        content_frame.paragraphs[0].font.size = Pt(18)
        content_frame.paragraphs[0].font.color.rgb = RGBColor(51, 51, 51)

        # Slide 3: Conclusiones (si existe)
        if "conclusiones" in data:
            slide3 = prs.slides.add_slide(prs.slide_layouts[6])
            title_box3 = slide3.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
            title_frame3 = title_box3.text_frame
            title_frame3.text = "Conclusiones"
            title_frame3.paragraphs[0].font.size = Pt(40)
            title_frame3.paragraphs[0].font.bold = True
            title_frame3.paragraphs[0].font.color.rgb = RGBColor(*primary_color)

            conclusion_box = slide3.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
            conclusion_frame = conclusion_box.text_frame
            conclusion_frame.text = data.get("conclusiones", "")
            conclusion_frame.paragraphs[0].font.size = Pt(18)

        # Slide final: Footer
        slide_final = prs.slides.add_slide(prs.slide_layouts[6])
        background_final = slide_final.background
        fill_final = background_final.fill
        fill_final.solid()
        fill_final.fore_color.rgb = RGBColor(*secondary_color)

        footer_box = slide_final.shapes.add_textbox(Inches(0.5), Inches(3), Inches(9), Inches(1.5))
        footer_frame = footer_box.text_frame
        footer_frame.text = self.theme.get("footer_text", "Documento Profesional")
        footer_frame.paragraphs[0].font.size = Pt(32)
        footer_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        footer_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Guardar
        prs.save(str(output_path))
        logger.info(f"PowerPoint document saved: {output_path}")

    @staticmethod
    def _hex_to_rgb(hex_color: str) -> tuple:
        """Convierte color hex a RGB"""
        hex_color = hex_color.lstrip("#")
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
