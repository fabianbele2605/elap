"""Word (.docx) Generator usando docxtpl + python-docx para items dinámicos"""

import logging
from pathlib import Path
from typing import Dict, Any
from datetime import datetime
from copy import deepcopy
import xml.etree.ElementTree as ET

logger = logging.getLogger(__name__)

try:
    from docxtpl import DocxTemplate
except ImportError:
    logger.warning("docxtpl no instalado. Instala con: pip install docxtpl")
    DocxTemplate = None

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
except ImportError:
    logger.warning("python-docx no instalado. Instala con: pip install python-docx")
    Document = None


class WordGenerator:
    """Genera documentos Word profesionales con templates"""

    def __init__(self, theme: Dict[str, Any], language: str = "es"):
        """
        Args:
            theme: Configuración del tema corporativo
            language: Idioma (es, en)
        """
        self.theme = theme
        self.language = language
        self.templates_dir = Path(__file__).parent.parent / "templates" / "word"

    def generate(self, document_type: str, data: Dict[str, Any], output_path: Path) -> None:
        """
        Genera documento Word usando docxtpl para variables simples
        y python-docx para items dinámicos en tablas

        Args:
            document_type: Tipo de documento (contract, report, invoice)
            data: Datos para rellenar template
            output_path: Ruta de salida
        """
        if not DocxTemplate:
            logger.error("docxtpl no instalado. Instala con: pip install docxtpl")
            raise ImportError("docxtpl required")

        # Obtener template según tipo de documento
        template_path = self._get_template_path(document_type)

        if not template_path.exists():
            logger.warning(f"Template no encontrado: {template_path}. Creando default...")
            self._create_default_template(template_path, document_type)

        try:
            # Crear directorio si no existe
            output_path.parent.mkdir(parents=True, exist_ok=True)

            # Paso 1: Cargar y renderizar variables simples con docxtpl
            doc = DocxTemplate(str(template_path))
            context = self._prepare_context(data, document_type)

            # Para facturas: NO pasar items a docxtpl (será manejado por python-docx)
            items_backup = None
            if document_type == "invoice" and "items" in context:
                items_backup = context.pop("items")

            doc.render(context)

            # Paso 2: Manejo especial de items dinámicos (para factura)
            if document_type == "invoice" and items_backup:
                # Guardar temporalmente
                temp_path = output_path.parent / f"temp_{output_path.name}"
                doc.save(str(temp_path))

                # Procesar items dinámicos con python-docx
                self._add_invoice_items_to_table(str(temp_path), items_backup, str(output_path))
                temp_path.unlink()  # Eliminar temporal
                logger.info(f"Word document generated with {len(items_backup)} items: {output_path}")
            else:
                # Guardar documento
                doc.save(str(output_path))
                logger.info(f"Word document generated: {output_path}")

        except Exception as e:
            logger.error(f"Error generating Word document: {e}")
            raise

    def _add_invoice_items_to_table(self, doc_path: str, items: list, output_path: str) -> None:
        """Agrega items dinámicos a tabla de factura clonando filas"""
        try:
            if not Document:
                logger.error("python-docx no disponible")
                return

            doc = Document(doc_path)

            # Encontrar tabla de items (segunda tabla típicamente)
            if len(doc.tables) < 2:
                logger.warning("No se encontró tabla de items, guardando sin procesar")
                doc.save(output_path)
                return

            items_table = doc.tables[1]

            # Encontrar índice de fila template
            template_row_idx = None
            for idx, row in enumerate(items_table.rows):
                row_text = " ".join([cell.text for cell in row.cells])
                if "{{item.descripcion}}" in row_text or "{{item.precio_unitario}}" in row_text:
                    template_row_idx = idx
                    break

            if template_row_idx is None:
                logger.warning("No se encontró fila template, guardando sin procesar items")
                doc.save(output_path)
                return

            # Obtener acceso a XML de la tabla
            tbl = items_table._element
            template_row_element = items_table.rows[template_row_idx]._element

            # Clonar fila para cada item
            for item_data in items:
                # Crear copia del elemento fila
                new_row_element = deepcopy(template_row_element)

                # Reemplazar variables en la fila clonada
                for text_elem in new_row_element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                    if text_elem.text:
                        if "{{item.descripcion}}" in text_elem.text:
                            text_elem.text = str(item_data.get("descripcion", ""))
                        elif "{{item.precio_unitario}}" in text_elem.text:
                            text_elem.text = str(item_data.get("precio_unitario", ""))
                        elif "{{item.cantidad}}" in text_elem.text:
                            text_elem.text = str(item_data.get("cantidad", ""))
                        elif "{{item.subtotal}}" in text_elem.text:
                            text_elem.text = str(item_data.get("subtotal", ""))

                # Insertar nueva fila en tabla
                tbl.insert(tbl.index(template_row_element) + 1, new_row_element)

            # Eliminar fila template original
            tbl.remove(template_row_element)

            # Guardar documento
            doc.save(output_path)
            logger.info(f"Added {len(items)} items to invoice table")

        except Exception as e:
            logger.error(f"Error adding invoice items: {e}")
            # Fallback: guardar el documento como está
            try:
                doc = Document(doc_path)
                doc.save(output_path)
            except:
                pass

    def _get_template_path(self, document_type: str) -> Path:
        """Obtiene ruta del template según tipo"""
        templates = {
            "contract": "contract_template.docx",
            "report": "report_template.docx",
            "invoice": "invoice_template.docx",
            "job_offer": "job_offer_template.docx",
            "hr_document": "hr_document_template.docx",
        }
        template_name = templates.get(document_type, "default_template.docx")
        return self.templates_dir / template_name

    def _prepare_context(self, data: Dict[str, Any], document_type: str) -> Dict[str, Any]:
        """Prepara contexto para renderizar template con docxtpl"""
        # Contexto base
        context = {
            # Datos del tema corporativo
            "empresa_nombre": self.theme.get("name", "Empresa"),
            "footer_text": self.theme.get("footer_text", ""),
            "primary_color": self.theme.get("primary_color", "#003366"),

            # Fechas
            "fecha_generacion": datetime.now().strftime("%d de %B de %Y"),
            "fecha_hoy": datetime.now().strftime("%Y-%m-%d"),
            "fecha_hoy_corta": datetime.now().strftime("%d/%m/%Y"),

            # Incluir todos los datos del usuario
            **data
        }

        # Procesamiento específico según tipo de documento
        if document_type == "contract":
            context = self._prepare_contract_context(context)
        elif document_type == "report":
            context = self._prepare_report_context(context)
        elif document_type == "invoice":
            context = self._prepare_invoice_context(context)

        return context

    def _prepare_contract_context(self, context: Dict) -> Dict:
        """Prepara contexto para contrato"""
        if "beneficios" not in context:
            context["beneficios"] = []
        if "responsabilidades" not in context:
            context["responsabilidades"] = []
        return context

    def _prepare_report_context(self, context: Dict) -> Dict:
        """Prepara contexto para reporte"""
        if "metricas" not in context:
            context["metricas"] = {}
        return context

    def _prepare_invoice_context(self, context: Dict) -> Dict:
        """Prepara contexto para factura"""
        # Calcular totales de items
        if "items" in context and isinstance(context["items"], list):
            for item in context["items"]:
                cantidad = float(item.get("cantidad", 1))
                precio = float(item.get("precio_unitario", 0))
                subtotal_item = cantidad * precio
                item["subtotal"] = f"${subtotal_item:,.0f}"
                item["precio_unitario"] = f"${precio:,.0f}"
                item["cantidad"] = int(cantidad)

            # Calcular totales
            subtotal = sum(float(item.get("precio_unitario", "0").replace("$", "").replace(",", "")) * item.get("cantidad", 1) for item in context["items"])
            context["subtotal"] = f"${subtotal:,.0f}"
            context["impuesto"] = f"${subtotal * 0.19:,.0f}"
            context["total"] = f"${subtotal * 1.19:,.0f}"
        else:
            context["subtotal"] = "$0"
            context["impuesto"] = "$0"
            context["total"] = "$0"
            context["items"] = []

        return context

    def _create_default_template(self, template_path: Path, document_type: str) -> None:
        """Crea template por defecto si no existe"""
        if not Document:
            logger.error("python-docx no disponible")
            return

        template_path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()

        title = doc.add_paragraph()
        title_run = title.add_run(f"Template: {document_type.upper()}")
        title_run.font.size = Pt(16)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(45, 80, 22)

        doc.add_paragraph("Este es un template por defecto. Personaliza según necesites.")

        if document_type == "contract":
            doc.add_paragraph("Empresa: {{empresa}}")
            doc.add_paragraph("Empleado: {{empleado}}")
            doc.add_paragraph("Cargo: {{cargo}}")

        doc.save(str(template_path))
        logger.info(f"Default template created: {template_path}")
