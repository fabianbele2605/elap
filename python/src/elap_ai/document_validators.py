"""
Validadores automáticos para documentos profesionales
Implementa CAPA 4: QA AUTOMÁTICO del Document Quality Standard
"""

import logging
import re
from typing import Dict, List, Tuple
from pathlib import Path
from datetime import datetime

logger = logging.getLogger(__name__)


class DocumentValidator:
    """Validador universal para documentos profesionales"""

    # Palabras clave prohibidas (placeholders, TODO, etc)
    FORBIDDEN_KEYWORDS = [
        r'\bTODO\b',
        r'\[insertar',
        r'\[TODO\]',
        r'\bxxx\b',
        r'lorem ipsum',
        r'\bLorem\b',
        r'undefined',
        r'placeholder',
        r'\[completar\]',
        r'\[falta\]',
    ]

    def __init__(self):
        self.issues = []
        self.warnings = []
        self.checks_passed = 0
        self.checks_total = 0

    def validate_content(self, content: str, filename: str) -> Dict:
        """Valida contenido del documento (busca placeholders, etc)"""
        self.issues = []
        self.checks_total = 0
        self.checks_passed = 0

        # Check 1: Buscar palabras prohibidas
        self.checks_total += 1
        forbidden_found = self._check_forbidden_keywords(content)
        if not forbidden_found:
            self.checks_passed += 1
        else:
            for keyword, line_num in forbidden_found:
                self.issues.append(f"Placeholder encontrado: '{keyword}' (potencial línea {line_num})")

        # Check 2: Buscar valores None/undefined
        self.checks_total += 1
        none_values = self._check_none_values(content)
        if not none_values:
            self.checks_passed += 1
        else:
            self.issues.append(f"Detectados {len(none_values)} valores 'None' o 'undefined'")

        # Check 3: Validar que no sea contenido vacío
        self.checks_total += 1
        if content.strip() and len(content.strip()) > 50:
            self.checks_passed += 1
        else:
            self.issues.append("Contenido muy corto o vacío")

        return {
            "check_type": "content_validation",
            "status": "✅ PASS" if not self.issues else "❌ FAIL",
            "issues": self.issues,
            "score": round((self.checks_passed / self.checks_total * 100)) if self.checks_total > 0 else 0
        }

    def validate_pdf(self, filepath: str) -> Dict:
        """Valida PDF (estructura, texto, etc)"""
        try:
            import pdfplumber
        except ImportError:
            logger.warning("pdfplumber no instalado, skipping PDF validation")
            return {"status": "SKIP", "reason": "pdfplumber not installed"}

        self.issues = []
        self.checks_total = 0
        self.checks_passed = 0

        try:
            with pdfplumber.open(filepath) as pdf:
                # Check 1: PDF tiene contenido
                self.checks_total += 1
                if len(pdf.pages) > 0:
                    self.checks_passed += 1
                else:
                    self.issues.append("PDF vacío (0 páginas)")

                # Check 2: Texto es seleccionable (no solo imágenes)
                self.checks_total += 1
                total_text = ""
                for page in pdf.pages:
                    text = page.extract_text() or ""
                    total_text += text

                if len(total_text.strip()) > 100:
                    self.checks_passed += 1
                else:
                    self.warnings.append("PDF puede contener solo imágenes (texto no seleccionable)")

                # Check 3: Detectar placeholders en texto
                self.checks_total += 1
                forbidden = self._check_forbidden_keywords(total_text)
                if not forbidden:
                    self.checks_passed += 1
                else:
                    for kw, _ in forbidden:
                        self.issues.append(f"Placeholder en PDF: '{kw}'")

                # Check 4: Verificar números de página
                self.checks_total += 1
                if len(pdf.pages) > 1:
                    # Si es multi-página, debería tener números
                    last_page_text = pdf.pages[-1].extract_text() or ""
                    if "página" in last_page_text.lower() or re.search(r'\d+\s*de\s*\d+', last_page_text):
                        self.checks_passed += 1
                    else:
                        self.warnings.append("PDF no detecta números de página")
                else:
                    self.checks_passed += 1

        except Exception as e:
            return {
                "check_type": "pdf_validation",
                "status": "❌ ERROR",
                "error": str(e)
            }

        return {
            "check_type": "pdf_validation",
            "status": "✅ PASS" if not self.issues else "⚠️ WARN" if not self.issues else "❌ FAIL",
            "issues": self.issues,
            "warnings": self.warnings,
            "pages": len(pdf.pages) if 'pdf' in locals() else 0,
            "score": round((self.checks_passed / self.checks_total * 100)) if self.checks_total > 0 else 0
        }

    def validate_excel(self, filepath: str) -> Dict:
        """Valida Excel (fórmulas, formato, etc)"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {"status": "SKIP", "reason": "openpyxl not installed"}

        self.issues = []
        self.warnings = []
        self.checks_total = 0
        self.checks_passed = 0

        try:
            wb = load_workbook(filepath, data_only=False)

            # Check 1: Tiene hojas
            self.checks_total += 1
            if len(wb.sheetnames) > 0:
                self.checks_passed += 1
            else:
                self.issues.append("Excel sin hojas")

            # Check 2: Detectar fórmulas con errores
            self.checks_total += 1
            error_formulas = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                            # Es una fórmula
                            if any(err in cell.value for err in ['#DIV/0', '#REF', '#N/A', '#VALUE']):
                                error_formulas.append(f"{sheet}!{cell.coordinate}")

            if not error_formulas:
                self.checks_passed += 1
            else:
                self.issues.append(f"Fórmulas con errores: {', '.join(error_formulas[:5])}")

            # Check 3: Hojas organizadas
            self.checks_total += 1
            if any(x in wb.sheetnames for x in ['Resumen', 'Datos', 'Gráficas', 'Análisis']):
                self.checks_passed += 1
            else:
                self.warnings.append("Estructura de hojas puede mejorarse")

            # Check 4: Detectar placeholders
            self.checks_total += 1
            has_placeholders = False
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str):
                            if self._has_forbidden_keyword(cell.value):
                                has_placeholders = True
                                self.issues.append(f"Placeholder en {sheet}!{cell.coordinate}: {cell.value[:50]}")

            if not has_placeholders:
                self.checks_passed += 1

        except Exception as e:
            return {
                "check_type": "excel_validation",
                "status": "❌ ERROR",
                "error": str(e)
            }

        return {
            "check_type": "excel_validation",
            "status": "✅ PASS" if not self.issues else "⚠️ WARN" if self.warnings and not self.issues else "❌ FAIL",
            "issues": self.issues,
            "warnings": self.warnings,
            "sheet_count": len(wb.sheetnames),
            "sheets": wb.sheetnames,
            "score": round((self.checks_passed / self.checks_total * 100)) if self.checks_total > 0 else 0
        }

    def validate_word(self, filepath: str) -> Dict:
        """Valida Word (estructura, estilos, etc)"""
        try:
            from docx import Document
        except ImportError:
            return {"status": "SKIP", "reason": "python-docx not installed"}

        self.issues = []
        self.warnings = []
        self.checks_total = 0
        self.checks_passed = 0

        try:
            doc = Document(filepath)

            # Check 1: Tiene contenido
            self.checks_total += 1
            if len(doc.paragraphs) > 0:
                self.checks_passed += 1
            else:
                self.issues.append("Word documento vacío")

            # Check 2: Detectar placeholders
            self.checks_total += 1
            has_placeholders = False
            for para in doc.paragraphs:
                if self._has_forbidden_keyword(para.text):
                    has_placeholders = True
                    self.issues.append(f"Placeholder encontrado: {para.text[:60]}")
                    break

            if not has_placeholders:
                self.checks_passed += 1

            # Check 3: Verificar tabla de contenidos
            self.checks_total += 1
            has_toc = any('TOC' in str(para._element.xml) for para in doc.paragraphs)
            if has_toc or any('Tabla de' in para.text for para in doc.paragraphs):
                self.checks_passed += 1
            else:
                self.warnings.append("No se detecta tabla de contenidos")

            # Check 4: Estilos consistentes
            self.checks_total += 1
            styles_used = [para.style.name for para in doc.paragraphs if para.style]
            if len(set(styles_used)) <= 10:  # Máximo 10 estilos diferentes
                self.checks_passed += 1
            else:
                self.warnings.append(f"Demasiados estilos diferentes ({len(set(styles_used))})")

        except Exception as e:
            return {
                "check_type": "word_validation",
                "status": "❌ ERROR",
                "error": str(e)
            }

        return {
            "check_type": "word_validation",
            "status": "✅ PASS" if not self.issues else "⚠️ WARN" if self.warnings and not self.issues else "❌ FAIL",
            "issues": self.issues,
            "warnings": self.warnings,
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "score": round((self.checks_passed / self.checks_total * 100)) if self.checks_total > 0 else 0
        }

    def generate_qa_report(self, filepath: str, agent_name: str, document_type: str) -> Dict:
        """Genera reporte completo de QA"""
        file_ext = Path(filepath).suffix.lower()

        # Validar contenido primero
        with open(filepath, 'rb') as f:
            content = f.read().decode('utf-8', errors='ignore')

        content_check = self.validate_content(content, Path(filepath).name)

        # Validar según tipo
        if document_type == 'pdf' or file_ext == '.pdf':
            format_check = self.validate_pdf(filepath)
        elif document_type == 'excel' or file_ext == '.xlsx':
            format_check = self.validate_excel(filepath)
        elif document_type == 'word' or file_ext == '.docx':
            format_check = self.validate_word(filepath)
        else:
            format_check = {"status": "UNKNOWN", "message": f"Tipo {document_type} no soportado"}

        # Calcular score general
        scores = [content_check.get('score', 0), format_check.get('score', 0)]
        overall_score = sum(scores) // len(scores) if scores else 0

        # Determinar status
        has_critical_issues = bool(content_check.get('issues', []) + format_check.get('issues', []))
        status = "🔴 FAIL" if has_critical_issues else "🟡 WARN" if format_check.get('warnings', []) else "🟢 PASS"

        report = {
            "timestamp": datetime.now().isoformat(),
            "document": Path(filepath).name,
            "agent": agent_name,
            "document_type": document_type,
            "format": file_ext.replace('.', ''),
            "overall_score": overall_score,
            "status": status,
            "checks": {
                "content": content_check,
                "format": format_check
            },
            "all_issues": content_check.get('issues', []) + format_check.get('issues', []),
            "all_warnings": format_check.get('warnings', []),
            "ready_for_delivery": not has_critical_issues
        }

        return report

    # Métodos auxiliares privados

    def _check_forbidden_keywords(self, text: str) -> List[Tuple[str, int]]:
        """Busca palabras prohibidas en el texto"""
        found = []
        for keyword_pattern in self.FORBIDDEN_KEYWORDS:
            matches = re.finditer(keyword_pattern, text, re.IGNORECASE)
            for match in matches:
                # Estimar número de línea
                line_num = text[:match.start()].count('\n') + 1
                found.append((match.group(), line_num))
        return found

    def _check_none_values(self, text: str) -> List[str]:
        """Detecta valores 'None' y 'undefined'"""
        none_pattern = r'\b(None|undefined|null|NULL)\b'
        return re.findall(none_pattern, text)

    def _has_forbidden_keyword(self, text: str) -> bool:
        """Verifica si el texto contiene palabras prohibidas"""
        for pattern in self.FORBIDDEN_KEYWORDS:
            if re.search(pattern, text, re.IGNORECASE):
                return True
        return False
